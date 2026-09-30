"""Pairs of large documents for the stress tests: two versions of one document,
as Word files and as plain text, at several sizes (paragraphs).

The second version has 4% of paragraphs reworded, 1% removed, 1% added and a
block of 20 paragraphs moved, as a long document's revision might.

Run: python3 scripts/stress/make-docs.py [sizes …]   (default: 1000 5000 10000 25000 50000)
Writes tests/fixtures/out/stress/doc-<size>-a.docx, -b.docx, -a.txt and -b.txt.
"""
import os
import random
import sys

from docx import Document

OUT = os.path.join(os.path.dirname(__file__), '..', '..', 'tests', 'fixtures', 'out', 'stress')
WORDS = (
    'the contract party shall provide services under this agreement including delivery payment terms notice '
    'period liability insurance confidential information warranty scope schedule fees invoice approval client '
    'consultant each month within days written consent reasonable efforts subject to any other law court'
).split()


def sentence(rng):
    n = rng.randint(8, 30)
    return ' '.join(rng.choice(WORDS) for _ in range(n)).capitalize() + '.'


def versions(size, seed=7):
    rng = random.Random(seed)
    base = []
    for i in range(size):
        if i % 40 == 0:
            base.append(('h', f'Section {i // 40 + 1}'))
        else:
            base.append(('p', ' '.join(sentence(rng) for _ in range(rng.randint(1, 3)))))
    edited = []
    for kind, text in base:
        r = rng.random()
        if r < 0.01:
            continue  # removed
        if r < 0.05 and kind == 'p':
            ws = text.split(' ')
            for _ in range(3):
                ws[rng.randrange(len(ws))] = rng.choice(WORDS)
            text = ' '.join(ws)
        edited.append((kind, text))
        if rng.random() < 0.01:
            edited.append(('p', sentence(rng)))  # added
    # A block moved from the first third to the last.
    at = len(edited) // 3
    block = edited[at:at + 20]
    del edited[at:at + 20]
    edited[len(edited) * 3 // 4:len(edited) * 3 // 4] = block
    return base, edited


def write(items, stem):
    d = Document()
    for kind, text in items:
        if kind == 'h':
            d.add_heading(text, level=1)
        else:
            d.add_paragraph(text)
    d.save(stem + '.docx')
    with open(stem + '.txt', 'w') as f:
        f.write('\n\n'.join(text for _, text in items) + '\n')


def main():
    sizes = [int(s) for s in sys.argv[1:]] or [1000, 5000, 10000, 25000, 50000]
    os.makedirs(OUT, exist_ok=True)
    for size in sizes:
        a, b = versions(size)
        write(a, os.path.join(OUT, f'doc-{size}-a'))
        write(b, os.path.join(OUT, f'doc-{size}-b'))
        words = sum(len(t.split()) for _, t in a)
        print(f'{size} paragraphs: {words} words ({words // 500} pages at 500 words a page)')


if __name__ == '__main__':
    main()
