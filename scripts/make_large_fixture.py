"""Generate two ~3,000-paragraph documents for tests/unit/perf.bench.test.ts.

Run: python3 scripts/make_large_fixture.py tests/fixtures/out
"""
import random, sys
from docx import Document
random.seed(4)
words = "the contract party shall provide services under this agreement including delivery payment terms notice period liability insurance confidential information warranty scope schedule fees invoice approval client consultant".split()
def sentence(n=None):
    n = n or random.randint(8, 30)
    return ' '.join(random.choice(words) for _ in range(n)).capitalize() + '.'
base = []
for i in range(3000):
    if i % 40 == 0:
        base.append(('h', f'Section {i // 40 + 1}'))
    else:
        base.append(('p', ' '.join(sentence() for _ in range(random.randint(1, 4)))))
def build(items, path):
    d = Document()
    for kind, text in items:
        if kind == 'h': d.add_heading(text, level=1)
        else: d.add_paragraph(text)
    d.save(path)
build(base, sys.argv[1] + '/big-a.docx')
edited = []
for kind, text in base:
    r = random.random()
    if r < 0.02: continue  # removed
    if r < 0.06 and kind == 'p':
        ws = text.split(' ')
        for _ in range(3):
            ws[random.randrange(len(ws))] = random.choice(words)
        text = ' '.join(ws)
    edited.append((kind, text))
    if random.random() < 0.02: edited.append(('p', sentence()))
build(edited, sys.argv[1] + '/big-b.docx')
print('ok')
